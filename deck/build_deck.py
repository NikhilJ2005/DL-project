"""Builds the project presentation in the MITS Kochi review template style.

Run:  python deck/build_deck.py   ->  deck/Mood_Companion_FDL.pptx
All coordinates are given in template pixels (2000 x 1125 canvas, 16:9) and converted to EMU.
"""
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
OUT = os.path.join(HERE, "Mood_Companion_FDL.pptx")

RED = RGBColor(0xCC, 0x00, 0x00)
BLACK = RGBColor(0, 0, 0)
GREY_TXT = RGBColor(0x59, 0x59, 0x59)
CARD = RGBColor(0xF2, 0xF2, 0xF2)
ROW_ALT = RGBColor(0xF2, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE = RGBColor(0xF4, 0x6B, 0x22)
FONT = "Calibri"

PROJECT = "Mood Companion"
FOOTER = f"{PROJECT}, Dept of AI and DS"
W_EMU = 12192000  # 13.333 in
SCALE = W_EMU / 2000


def px(v):
    return Emu(int(v * SCALE))


prs = Presentation()
prs.slide_width, prs.slide_height = px(2000), px(1125)
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers
def text(slide, x, y, w, h, runs, size=20, color=BLACK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None):
    """runs: str, or list of paragraphs; each paragraph is a str or a list of (text, bold, color)."""
    box = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(45720)
    tf.margin_top = tf.margin_bottom = Emu(22860)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        parts = para if isinstance(para, list) else [(para, bold, color)]
        for t, b, c in parts:
            r = p.add_run()
            r.text = t
            r.font.name, r.font.size, r.font.bold = FONT, Pt(size), b
            r.font.color.rgb = c
    return box


def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=1.5, radius=None):
    s = slide.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    return s


def label_shape(s, txt, size=16, color=BLACK, bold=False):
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Emu(45720)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = txt.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold
        r.font.color.rgb = color
    return s


def box(slide, x, y, w, h, txt, fill=WHITE, line=BLACK, size=15, bold=False, color=BLACK,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = rect(slide, x, y, w, h, fill, line, shape, line_w=1.25, radius=0.15)
    return label_shape(s, txt, size, color, bold)


def arrow(slide, x1, y1, x2, y2, color=BLACK, dashed=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    c.line.color.rgb = color
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    if dashed:
        dash = etree.SubElement(ln, qn("a:prstDash"))
        dash.set("val", "dash")
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    return c


def logo(slide, big=False):
    path = os.path.join(ASSETS, "mits_logo_hd.png")
    if big:
        slide.shapes.add_picture(path, px(818), px(50), px(370), px(212))
    else:
        slide.shapes.add_picture(path, px(1712), px(32), px(218), px(124))


def footer(slide, n, show_text=True):
    rect(slide, 0, 1020, 2000, 105, fill=RED)
    if show_text:
        text(slide, 50, 1045, 900, 55, FOOTER, size=14, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 1800, 1045, 150, 55, str(n), size=16, color=WHITE, bold=True,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def title(slide, t, y=100):
    text(slide, 110, y, 1550, 110, t, size=40, bold=True)


def new_slide(n, heading=None, title_y=100):
    s = prs.slides.add_slide(BLANK)
    logo(s)
    footer(s, n)
    if heading:
        title(s, heading, title_y)
    return s


def table(slide, x, y, col_w, row_h, rows, size=16, header_size=18, first_col_bold=True):
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, px(x), px(y), px(sum(col_w)), px(row_h * n_rows))
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:  # plain "No style, table grid"
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"
    for j, w in enumerate(col_w):
        tbl.columns[j].width = px(w)
    for i in range(n_rows):
        tbl.rows[i].height = px(row_h)
        for j in range(n_cols):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Emu(54864)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = RED
            else:
                cell.fill.fore_color.rgb = ROW_ALT if i % 2 == 1 else WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = rows[i][j]
            r.font.name = FONT
            r.font.size = Pt(header_size if i == 0 else size)
            r.font.bold = i == 0 or (j == 0 and first_col_bold)
            r.font.color.rgb = WHITE if i == 0 else BLACK
    return tbl


# ---------------------------------------------------------------- 1. Title
s = prs.slides.add_slide(BLANK)
logo(s, big=True)
footer(s, 1, show_text=False)
text(s, 100, 285, 1800, 80, PROJECT, size=44, bold=True, align=PP_ALIGN.CENTER)
text(s, 100, 375, 1800, 150,
     ["Deep Sentiment-Conditioned Neural Contextual Bandits", "for Real-Time Interactive Desktop Companions"],
     size=34, bold=True, align=PP_ALIGN.CENTER)
text(s, 100, 560, 1800, 70, "Foundations of Deep Learning", size=30, align=PP_ALIGN.CENTER)
text(s, 100, 650, 1800, 60, "Academic Year 2026–27  |  31 October 2026", size=24, bold=True,
     align=PP_ALIGN.CENTER)
text(s, 1250, 760, 650, 255,
     [[("Team Members", True, BLACK)], "Reuben Sabu V (MUT23CA060)", "Sidharth Ravi (MUT23CA063)",
      "Jino Jenz (MUT23CA035)", "Nikhil Jimmy (MUT23CA055)"],
     size=20, align=PP_ALIGN.RIGHT)

# ---------------------------------------------------------------- 2. Problem Statement
s = new_slide(2, "Problem Statement")
text(s, 80, 225, 1840, 500,
     "Desktop companions and chat assistants reply to every user the same way, whatever their mood. "
     "Fixed, hand-written rules cannot tell whether a stressed user wants a joke, a break or a "
     "breathing exercise, and they never improve from feedback. This project builds a lightweight, "
     "real-time companion that reads the user's emotion from text with a deep neural network and "
     "uses a neural contextual bandit to choose the response style most likely to help, learning "
     "from each thumbs-up / thumbs-down so that it personalises itself to the user while it runs.",
     size=24, color=RGBColor(0x26, 0x26, 0x26), align=PP_ALIGN.JUSTIFY, spacing=6)

# ---------------------------------------------------------------- 3. Objectives
s = new_slide(3, "Objectives")
objectives = [
    ("Detect emotion:", " Train a deep MLP on TF-IDF features to classify text into 6 emotions."),
    ("Build context:", " Combine emotion probabilities, time of day, session length and last reward into a 10-d context."),
    ("Decide:", " Use a neural contextual bandit (MLP + ε-greedy) to pick 1 of 8 companion actions."),
    ("Learn online:", " Update the bandit in real time from user thumbs-up / thumbs-down rewards."),
    ("Evaluate & demo:", " Compare against random, fixed and oracle policies; deliver an interactive Streamlit app."),
]
for i, (head, body) in enumerate(objectives):
    y = 250 + i * 144
    c = rect(s, 120, y, 104, 104, fill=RED, shape=MSO_SHAPE.OVAL)
    label_shape(c, f"O{i + 1}", size=20, color=WHITE, bold=True)
    text(s, 255, y + 2, 1400, 100, [[(head, True, BLACK), (body, False, GREY_TXT)]], size=22,
         anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 4. Project Overview
s = new_slide(4, "Project Overview")
cards = [
    ("Problem", "Assistants give one-size-fits-all replies and cannot adapt to a user's mood."),
    ("Motivation", "The wrong reply at a stressful moment annoys users; the right one supports them."),
    ("Need", "A model that senses emotion and learns which response works, from live feedback."),
    ("Target Users", "Students, remote workers and anyone using a desktop assistant for long sessions."),
    ("Scope", "Text emotion model, neural bandit, simulator and a 4-page Streamlit web app."),
    ("Expected Impact", "Responses that fit the user's mood and improve with every interaction."),
]
for i, (h, b) in enumerate(cards):
    x, y = 120 + (i % 3) * 430, 250 + (i // 3) * 370
    rect(s, x, y, 410, 345, fill=CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    text(s, x + 20, y + 30, 370, 60, h, size=21, bold=True, color=RGBColor(0x26, 0x26, 0x26))
    text(s, x + 20, y + 100, 370, 235, b, size=18, color=GREY_TXT)
rect(s, 1430, 250, 450, 230, fill=ORANGE)
text(s, 1430, 275, 450, 90, "SDG 3", size=44, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
text(s, 1430, 375, 450, 90, "Good Health and Well-being", size=20, color=WHITE, align=PP_ALIGN.CENTER)
rect(s, 1430, 500, 450, 465, line=ORANGE, line_w=2)
text(s, 1440, 510, 430, 450,
     [[("Target 3.4", True, BLACK)],
      [("Promote mental health and well-being.", False, GREY_TXT)],
      [("Mood Companion notices stress, sadness or anger and responds with breaks, breathing "
        "exercises or support — adapting to what helps each user.", False, GREY_TXT)]],
     size=17, spacing=4)

# ---------------------------------------------------------------- 5. Progress since First Review
s = new_slide(5, "Progress since First Review", title_y=60)
table(s, 36, 158, [350, 575, 1015], 75, [
    ["Parameter", "Initial Plan", "Current Status (31 Oct 2026)"],
    ["Problem / Scope", "Mood-aware companion idea", "Sentiment-conditioned neural bandit companion"],
    ["Dataset", "Emotion text data from Kaggle", "Kaggle Emotions dataset (6 labels) via kagglehub"],
    ["Emotion Model", "TF-IDF + small neural network", "TF-IDF → MLP (128-64-6) trained and saved"],
    ["Decision Model", "Contextual bandit for responses", "Neural bandit (10→64→64→8) with ε-greedy"],
    ["Evaluation", "Simulated user planned", "Done: 70.4% reward vs 40.6% for random"],
    ["Platform", "Streamlit web app", "4-page app built and browser-tested"],
], size=19, header_size=21)
table(s, 36, 735, [727, 1213], 85, [
    ["Key Change since Review 1", "Outcome"],
    ["1. Kaggle access may be blocked offline", "Added built-in sample data so the app always runs"],
    ["2. Context made sentiment-conditioned", "Bandit learns a different best action per mood"],
], size=19, header_size=21, first_col_bold=False)

# ---------------------------------------------------------------- 6. System Requirements and Design
s = new_slide(6, "System Requirements and Design")
layers = [
    ("Presentation", RGBColor(0xE8, 0xF0, 0xFB), ["Dataset &\nTraining page", "Live\nCompanion",
                                                   "Bandit\nSimulation", "How it\nWorks"]),
    ("Model\n(PyTorch)", RGBColor(0xE6, 0xF4, 0xEA), ["Emotion model\nTF-IDF → MLP", "Context\nbuilder (10-d)",
                                                       "Neural bandit\nMLP + ε-greedy", "Response\ntemplates"]),
    ("Learning", RGBColor(0xFD, 0xF1, 0xE0), ["Reward\n(up/down)", "Replay\nbuffer",
                                               "Online update\n(Adam, MSE)", "Simulated\nuser"]),
    ("Data &\nStorage", RGBColor(0xF1, 0xE8, 0xF8), ["Kaggle dataset\n(kagglehub)", "Offline\nsample CSV",
                                                     "models/\n*.pt weights", "Streamlit\nsession state"]),
]
for i, (name, colour, items) in enumerate(layers):
    y = 215 + i * 162
    text(s, 60, y + 25, 220, 90, name, size=14, bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 285, y, 935, 140, fill=colour, line=RGBColor(0x8C, 0x8C, 0x8C), shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         radius=0.08, line_w=1)
    for j, it in enumerate(items):
        box(s, 305 + j * 228, y + 18, 210, 104, it, size=12)
    if i < len(layers) - 1:
        arrow(s, 752, y + 140, 752, y + 162)
text(s, 285, 860, 935, 40, "(a) Overall system architecture", size=16, align=PP_ALIGN.CENTER)
modules = ["User Message", "Emotion Model", "Context Builder", "Neural Bandit", "Companion Response"]
for i, m in enumerate(modules):
    y = 230 + i * 128
    box(s, 1250, y, 390, 78, m, size=18, bold=True, shape=MSO_SHAPE.RECTANGLE)
    if i < len(modules) - 1:
        arrow(s, 1445, y + 78, 1445, y + 128)
box(s, 1700, 594, 200, 100, "Feedback\nLearner", size=18, bold=True, shape=MSO_SHAPE.RECTANGLE)
arrow(s, 1800, 822, 1800, 694)
arrow(s, 1640, 822, 1800, 822, )
arrow(s, 1700, 644, 1640, 644)
text(s, 1250, 860, 650, 40, "(b) Simplified module block diagram", size=16, align=PP_ALIGN.CENTER)
rect(s, 100, 910, 1800, 100, fill=CARD)
text(s, 110, 915, 1780, 90,
     [[("Software: ", True, BLACK), ("Python 3.10+, PyTorch, scikit-learn, Streamlit, kagglehub, pandas.   ", False, BLACK),
       ("Hardware: ", True, BLACK), ("any 4-core CPU, 8 GB RAM; no GPU needed (training ≈ 1 min).", False, BLACK)]],
     size=18, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 7. Methodology
s = new_slide(7, "Methodology", title_y=45)
fc = [  # (label, y, h, shape, fill)
    ("Start", 170, 50, MSO_SHAPE.ROUNDED_RECTANGLE, WHITE),
    ("User types a message", 250, 60, MSO_SHAPE.RECTANGLE, WHITE),
    ("Emotion MLP → 6 probs", 340, 60, MSO_SHAPE.RECTANGLE, WHITE),
    ("Build 10-d context", 430, 60, MSO_SHAPE.RECTANGLE, WHITE),
    ("rand < ε ?", 515, 100, MSO_SHAPE.DIAMOND, WHITE),
    ("Show response", 660, 60, MSO_SHAPE.RECTANGLE, RGBColor(0xE6, 0xF4, 0xEA)),
    ("Thumbs up / down = reward", 750, 60, MSO_SHAPE.RECTANGLE, RGBColor(0xE8, 0xF0, 0xFB)),
    ("Update bandit (MSE, decay ε)", 840, 60, MSO_SHAPE.RECTANGLE, RGBColor(0xF1, 0xE8, 0xF8)),
]
cx = 420
for i, (lbl, y, h, shp, fill) in enumerate(fc):
    w = 240 if shp == MSO_SHAPE.DIAMOND else 380
    box(s, cx - w / 2, y, w, h, lbl, fill=fill, size=14, shape=shp)
    if i < len(fc) - 1 and shp != MSO_SHAPE.DIAMOND:
        arrow(s, cx, y + h, cx, fc[i + 1][1])
# diamond branches
box(s, 40, 525, 190, 80, "Explore:\nrandom action", size=13, fill=RGBColor(0xFD, 0xF1, 0xE0))
box(s, 610, 525, 190, 80, "Exploit: best\npredicted action", size=13, fill=RGBColor(0xFD, 0xF1, 0xE0))
arrow(s, cx - 120, 565, 230, 565)
arrow(s, cx + 120, 565, 610, 565)
text(s, 240, 525, 60, 30, "yes", size=12)
text(s, 555, 525, 60, 30, "no", size=12)
arrow(s, 135, 605, cx - 80, 660)
arrow(s, 705, 605, cx + 80, 660)
# loop back to the next message
arrow(s, cx + 190, 870, 850, 870)
arrow(s, 850, 870, 850, 280)
arrow(s, 850, 280, cx + 190, 280)
text(s, 856, 640, 130, 60, "next\nmessage", size=11, color=GREY_TXT)
text(s, 100, 935, 640, 40, "Flowchart of the Sense → Decide → Learn loop", size=15, align=PP_ALIGN.CENTER)
steps = [
    ("1. Data:", " Kaggle emotions dataset via kagglehub."),
    ("2. Features:", " TF-IDF, 5,000 uni- and bi-grams."),
    ("3. Emotion MLP:", " 5000→128→64→6, ReLU, Dropout 0.3."),
    ("4. Context:", " emotion probs + time + session + last reward."),
    ("5. Reward net:", " MLP 10→64→64→8, one output per action."),
    ("6. ε-greedy:", " explore with prob. ε (0.3 → 0.05)."),
    ("7. Online learning:", " replay buffer, 4 Adam steps, MSE."),
    ("8. Evaluation:", " vs random, fixed and oracle policies."),
]
text(s, 980, 170, 990, 50, "Processing steps & workflow", size=24, bold=True)
text(s, 980, 220, 990, 400, [[(h, True, BLACK), (b, False, GREY_TXT)] for h, b in steps], size=18)
text(s, 980, 640, 990, 50, "AI-project items", size=24, bold=True)
items = [
    ("Dataset:", " ~20k sentences, 6 emotions (16k/2k/2k)."),
    ("Pre-processing:", " label filtering, TF-IDF."),
    ("Training:", " 8 epochs, batch 64, Adam lr 1e-3, CPU."),
    ("Testing:", " accuracy, confusion matrix, simulation."),
    ("Deployment:", " Streamlit web app."),
]
text(s, 1000, 690, 970, 320, [[("•  " + h, True, BLACK), (b, False, GREY_TXT)] for h, b in items], size=18)

# ---------------------------------------------------------------- 8. Technology Used
s = new_slide(8, "Technology Used")
table(s, 120, 240, [345, 570, 835], 63, [
    ["Category", "Choice", "Why it was selected"],
    ["Programming language", "Python 3.10+", "Standard language for deep learning and data science"],
    ["Deep learning", "PyTorch 2.x", "Simple, flexible neural networks; runs well on CPU"],
    ["Feature extraction", "scikit-learn (TF-IDF)", "Fast, reliable text-to-vector conversion"],
    ["Web framework", "Streamlit", "Interactive Python web app with no front-end code"],
    ["Dataset access", "kagglehub", "One-line download of the Kaggle dataset"],
    ["Data handling", "pandas, NumPy", "Loading, cleaning and numeric processing"],
    ["Visualisation", "Streamlit charts, matplotlib", "Loss curves, confusion matrix, reward plots"],
    ["Testing", "Playwright (headless Chromium)", "Automated check that every app page works"],
    ["IDE & tools", "VS Code, Git / GitHub", "Shared team tooling and version control"],
    ["Hardware", "4-core CPU, 8 GB RAM; no GPU", "Small models train in about a minute on CPU"],
], size=17, header_size=19)

# ---------------------------------------------------------------- 9. Results
s = new_slide(9, "Results")
stats = [
    ("70.4%", "neural bandit\nthumbs-up rate*"),
    ("40.6%", "random-choice\nbaseline*"),
    ("86.5%", "oracle upper\nbound*"),
    ("3.5 ms", "per online update\non CPU"),
    ("6 → 8", "emotions mapped\nto actions"),
    ("4 / 4", "app pages pass\nbrowser test"),
]
for i, (big, small) in enumerate(stats):
    x, y = 60 + (i % 3) * 360, 190 + (i // 3) * 215
    rect(s, x, y, 335, 200, line=BLACK, line_w=1.5)
    text(s, x, y + 15, 335, 70, big, size=36, bold=True, color=RED, align=PP_ALIGN.CENTER)
    text(s, x, y + 95, 335, 95, small, size=18, color=GREY_TXT, align=PP_ALIGN.CENTER)
s.shapes.add_picture(os.path.join(ASSETS, "results_chart.png"), px(1220), px(185), px(640), px(416))
text(s, 1110, 600, 860, 50, "*Simulated user, 2,000 rounds, assuming accurate emotion detection",
     size=13, color=GREY_TXT, align=PP_ALIGN.CENTER)
table(s, 33, 660, [650, 565, 700], 70, [
    ["Performance indicator", "Expected", "Achieved"],
    ["Emotion classifier accuracy", "≥ 80% on test set", "[fill in after full Kaggle training]"],
    ["Bandit beats random policy", "Clearly higher reward", "Yes: 70.4% vs 40.6%"],
    ["Bandit beats fixed best action", "Higher reward", "Yes: 70.4% vs 64.3%"],
    ["Real-time learning", "< 100 ms per feedback", "~3.5 ms per update on CPU"],
], size=19, header_size=21)

# ---------------------------------------------------------------- 10. Remaining Work
s = new_slide(10, "Remaining Work", title_y=80)
table(s, 12, 193, [330, 1018, 597], 88, [
    ["Activity", "Deliverable", "% Completed"],
    ["Full training", "Train on full Kaggle data; report accuracy", "50%"],
    ["User study", "Real users rate responses to validate the bandit", "0%"],
    ["Better model", "Upgrade emotion model to a transformer (DistilBERT)", "0%"],
    ["Exploration", "Compare ε-greedy with NeuralUCB / Thompson sampling", "10%"],
    ["Desktop app", "Always-on-top desktop overlay version of the companion", "0%"],
    ["Documentation", "Project report, user guide and presentation", "60%"],
], size=20, header_size=22)
rect(s, 12, 830, 1953, 90, fill=RED)
text(s, 15, 830, 1900, 90, "Overall completion ≈ 70% of planned scope", size=24, bold=True, color=WHITE,
     anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 11. Updated Timelines
s = new_slide(11, "Updated Timelines")
cols = [
    ("Completed", ["Problem study & literature survey: Oct 2026", "Emotion MLP + neural bandit: Oct 2026",
                   "Streamlit app (4 pages): Oct 2026", "Simulation evaluation: Oct 2026"]),
    ("Ongoing", ["Full Kaggle training & accuracy: Nov 2026", "Exploration-strategy comparison: Nov 2026",
                 "Project report: Nov 2026"]),
    ("Pending", ["Real-user study: Nov–Dec 2026", "Transformer emotion model: Dec 2026",
                 "Final presentation & viva: Dec 2026"]),
]
for i, (h, items) in enumerate(cols):
    x = 120 + i * 600
    text(s, x, 250, 560, 60, h, size=24, bold=True)
    tb = text(s, x + 20, 330, 540, 650, items, size=22, color=RGBColor(0x26, 0x26, 0x26), spacing=10)
    for p in tb.text_frame.paragraphs:  # real bullets so they stay editable
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", "342900")
        pPr.set("indent", "-342900")
        bu = etree.SubElement(pPr, qn("a:buChar"))
        bu.set("char", "•")

# ---------------------------------------------------------------- 12. References
s = new_slide(12, "References")
refs = [
    'L. Li, W. Chu, J. Langford and R. E. Schapire, "A contextual-bandit approach to personalized news article recommendation," WWW, 2010.',
    'C. Riquelme, G. Tucker and J. Snoek, "Deep Bayesian bandits showdown," ICLR, 2018.',
    'D. Zhou, L. Li and Q. Gu, "Neural contextual bandits with UCB-based exploration," ICML, 2020.',
    'R. Allesiardo, R. Féraud and D. Bouneffouf, "A neural networks committee for the contextual bandit problem," ICONIP, 2014.',
    'E. Saravia et al., "CARER: Contextualized affect representations for emotion recognition," EMNLP, 2018.',
    'R. S. Sutton and A. G. Barto, Reinforcement Learning: An Introduction, 2nd ed., MIT Press, 2018.',
    'I. Goodfellow, Y. Bengio and A. Courville, Deep Learning, MIT Press, 2016.',
    'T. Lattimore and C. Szepesvári, Bandit Algorithms, Cambridge University Press, 2020.',
    'A. Paszke et al., "PyTorch: An imperative style, high-performance deep learning library," NeurIPS, 2019.',
    'F. Pedregosa et al., "Scikit-learn: Machine learning in Python," JMLR, 2011.',
    'P. Govi, "Emotions dataset for NLP," kaggle.com/datasets/praveengovi/emotions-dataset-for-nlp.',
    'Streamlit Inc., "Streamlit documentation," docs.streamlit.io, 2025.',
    'Kaggle, "kagglehub," github.com/Kaggle/kagglehub, 2025.',
]
half = 7
for col, chunk in enumerate([refs[:half], refs[half:]]):
    start = col * half
    text(s, 28 + col * 1000, 230, 950, 780,
         [[(f"[{start + k + 1}] ", True, BLACK), (r, False, BLACK)] for k, r in enumerate(chunk)],
         size=17, spacing=4)

# ---------------------------------------------------------------- 13. Thank You
s = new_slide(13)
text(s, 0, 370, 2000, 150, "Thank You", size=72, bold=True, color=RED, align=PP_ALIGN.CENTER)
text(s, 0, 570, 2000, 70, "Mood Companion: emotion → context → neural bandit → personalised response",
     size=28, color=GREY_TXT, align=PP_ALIGN.CENTER)
text(s, 0, 690, 2000, 60, "Questions?", size=26, align=PP_ALIGN.CENTER)

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
